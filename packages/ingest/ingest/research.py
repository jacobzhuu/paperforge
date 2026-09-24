"""Versioned, deterministic descriptive analysis. Never executes uploaded code."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import statistics
import zipfile
from collections import defaultdict
from typing import Any

VERSION = "descriptive-v1"
MAX_ROWS = 100000
MAX_COLUMNS = 100
MAX_GROUPS = 100


def digest(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, ensure_ascii=False, allow_nan=False).encode()
    ).hexdigest()


def read_table(content: bytes, filename: str, sheet: str | None = None) -> dict:
    if len(content) > 32 * 1024 * 1024:
        raise ValueError("文件超过 32 MiB")
    sheets = []
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 128 * 1024 * 1024:
                raise ValueError("工作簿解压后过大")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            sheets = book.sheetnames
            if sheet is None and len(sheets) != 1:
                return {"sheets": sheets, "needs_sheet": True}
            selected = sheet or sheets[0]
            if selected not in sheets:
                raise ValueError("工作表不存在")
            rows = []
            for values in book[selected].iter_rows():
                if len(rows) > MAX_ROWS or len(values) > MAX_COLUMNS:
                    raise ValueError("超过分析上限：100000 行 / 100 列")
                if any(c.data_type == "f" for c in values):
                    raise ValueError("请将公式转为可核对的数值后上传；分析不会执行工作簿公式")
                rows.append(["" if c.value is None else str(c.value).strip() for c in values])
        finally:
            book.close()
    elif filename.lower().endswith((".csv", ".tsv")):
        try:
            source = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("请上传 UTF-8 编码 CSV/TSV") from error
        rows = []
        for row in csv.reader(
            io.StringIO(source), delimiter="\t" if filename.lower().endswith(".tsv") else ","
        ):
            if len(rows) > MAX_ROWS or len(row) > MAX_COLUMNS:
                raise ValueError("超过分析上限：100000 行 / 100 列")
            rows.append([c.strip() for c in row])
    else:
        raise ValueError("分析仅支持 CSV、TSV、XLSX")
    rows = [r for r in rows if any(r)]
    if len(rows) < 2:
        raise ValueError("表格没有数据行")
    headers = rows[0]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ValueError("表头必须非空且不能重复")
    if any(len(r) != len(headers) for r in rows[1:]):
        raise ValueError("数据行列数与表头不一致")
    return {
        "headers": headers,
        "rows": rows[1:],
        "sheets": sheets,
        "sheet": sheet or (sheets[0] if sheets else None),
    }


def analyze(table: dict, spec: dict) -> dict:
    headers, rows = table["headers"], table["rows"]
    columns = spec.get("columns") or []
    group = spec.get("group_by")
    if not columns or len(columns) > 20 or len(set(columns)) != len(columns):
        raise ValueError("请选择 1–20 个不重复的数值字段")
    if any(c not in headers for c in columns) or (group and group not in headers):
        raise ValueError("分析字段不在表头中")
    if group in columns:
        raise ValueError("分组字段不能同时作为数值字段")
    units = spec.get("units") or {}
    if any(c not in units or not str(units[c]).strip() for c in columns):
        raise ValueError("请明确每个字段的单位（无量纲请填写 1）")
    groups: dict[str, list] = defaultdict(list)
    for row in rows:
        label = row[headers.index(group)] if group else "全部"
        if group and not label:
            raise ValueError("分组字段存在缺失值，请明确分组后重试")
        groups[label].append(row)
        if len(groups) > MAX_GROUPS:
            raise ValueError("分组超过 100 个，请选择分类字段")
    records = []
    for label, members in groups.items():
        for column in columns:
            values = []
            for row in members:
                raw = row[headers.index(column)]
                if raw.lower() in {"", "na", "n/a", "null"}:
                    continue
                try:
                    value = float(raw)
                except ValueError as error:
                    raise ValueError(f"{column} 存在非数值或混合单位；不会静默忽略") from error
                if not math.isfinite(value):
                    raise ValueError(f"{column} 包含非有限数值")
                values.append(value)
            record = {
                "group": label,
                "column": column,
                "unit": units[column],
                "n": len(values),
                "missing": len(members) - len(values),
                "mean": statistics.mean(values) if values else None,
                "median": statistics.median(values) if values else None,
                "min": min(values) if values else None,
                "max": max(values) if values else None,
                "std": statistics.stdev(values) if len(values) > 1 else None,
            }
            if any(isinstance(v, float) and not math.isfinite(v) for v in record.values()):
                raise ValueError("统计结果溢出，请调整输入量纲")
            records.append(record)
    return {
        "version": VERSION,
        "spec": spec,
        "row_count": len(rows),
        "records": records,
        "policy": {
            "std_ddof": 1,
            "missing": "exclude_per_column_and_report",
            "outliers": "retain",
            "inference": False,
        },
    }


def derived_table(result: dict) -> dict:
    headers = ["分组", "字段", "单位", "n", "missing", "mean", "median", "min", "max", "std"]
    rows, cells = [], {}
    for i, record in enumerate(result["records"]):
        row = [str(record[k]) for k in ("group", "column", "unit")]
        for key in headers[3:]:
            value = record[key]
            shown = (
                format(value, ".10g")
                if isinstance(value, float)
                else str(value)
                if value is not None
                else "不可计算"
            )
            row.append(shown)
            if value is not None:
                cells[f"{i}::{key}"] = shown
        rows.append(row)
    return {
        "type": "table",
        "headers": headers,
        "rows": rows,
        "numeric_cells": cells,
        "numbers": sorted(set(cells.values())),
        "row_count": len(rows),
        "column_count": len(headers),
        "analysis_version": VERSION,
    }


def chart_png(result: dict) -> bytes:
    import matplotlib

    matplotlib.use("Agg")
    matplotlib.rcParams["font.family"] = ["Noto Sans CJK SC", "DejaVu Sans"]
    matplotlib.rcParams["axes.unicode_minus"] = False
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure

    records = result["records"]
    columns = list(dict.fromkeys(r["column"] for r in records))
    figure = Figure(figsize=(8, max(3, len(columns) * 2.8)), layout="constrained")
    FigureCanvasAgg(figure)
    for idx, column in enumerate(columns):
        ax = figure.add_subplot(len(columns), 1, idx + 1)
        subset = [r for r in records if r["column"] == column and r["mean"] is not None]
        ax.bar(range(len(subset)), [r["mean"] for r in subset])
        ax.set_xticks(range(len(subset)), [r["group"] for r in subset], rotation=25)
        ax.set_ylabel(str(result["spec"]["units"][column]))
        ax.set_title(f"{column} — mean (descriptive)")
    output = io.BytesIO()
    figure.savefig(output, format="png", dpi=120)
    return output.getvalue()
